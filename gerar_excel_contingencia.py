# -*- coding: utf-8 -*-
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ARQUIVO_CSV = "contingencia_civel_api.csv"
ARQUIVO_EXCEL = "contingencia_civel_api.xlsx"

def gerar_excel():
    # Lê CSV
    dados = []
    with open(ARQUIVO_CSV, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        dados = list(reader)
    
    if not dados:
        print("ERRO: CSV vazio")
        return
    
    # Cria workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Contingência"
    
    # Headers
    headers = ["Processo", "Cliente", "Adverso", "Instância", "Advogado", "Advogado Adverso", "Assunto", "Natureza", "Status"]
    ws.append(headers)
    
    # Formatação header
    header_fill = PatternFill(start_color="233240", end_color="233240", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # Adiciona dados
    for row in dados:
        ws.append([
            row.get("Processo", ""),
            row.get("Cliente", ""),
            row.get("Adverso", ""),
            row.get("Instância", ""),
            row.get("Advogado", ""),
            row.get("Advogado Adverso", ""),
            row.get("Assunto", ""),
            row.get("Natureza", ""),
            row.get("Status", "")
        ])
    
    # Ajusta largura colunas
    ws.column_dimensions['A'].width = 22  # Processo
    ws.column_dimensions['B'].width = 25  # Cliente
    ws.column_dimensions['C'].width = 25  # Adverso
    ws.column_dimensions['D'].width = 18  # Instância
    ws.column_dimensions['E'].width = 25  # Advogado
    ws.column_dimensions['F'].width = 25  # Adv. Adverso
    ws.column_dimensions['G'].width = 35  # Assunto
    ws.column_dimensions['H'].width = 15  # Natureza
    ws.column_dimensions['I'].width = 18  # Status
    
    # Formatação dados
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    data_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = thin_border
            cell.alignment = data_alignment
    
    # Congela header
    ws.freeze_panes = "A2"
    
    # Salva
    wb.save(ARQUIVO_EXCEL)
    print(f"✅ Excel gerado: {ARQUIVO_EXCEL}")
    print(f"✅ Total de registros: {len(dados)}")

if __name__ == "__main__":
    gerar_excel()
